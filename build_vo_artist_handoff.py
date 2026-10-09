from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "VO_Artist_Handoff_Current_Active_VO_All_Scenes.docx"


WORD_SECTIONS = [
    ("Shloka Connected Words", [
        ("Vakratunda", "वक्रतुण्ड", "Va-kra-tun-da"),
        ("Mahakaya / Mahakaay", "महाकाय", "Ma-ha-ka-ya"),
        ("Suryakoti", "सूर्यकोटि", "Sur-ya-ko-ti"),
        ("Samaprabha", "समप्रभ", "Sa-ma-pra-bha"),
        ("Nirvighnam", "निर्विघ्नम्", "Nir-vigh-nam"),
        ("Kuru Me Deva", "कुरु मे देव", "kuru-me-deva"),
        ("Sarva Karyeshu", "सर्व कार्येषु", "sarva-kar-ye-shu"),
        ("Sarvada", "सर्वदा", "Sar-va-da"),
    ]),
    ("Names, Family, Culture", [
        ("Ganesha", "गणेश", "Ga-ne-sha"),
        ("Shiva Ji", "शिव जी", "Shi-va Ji"),
        ("Parvati Mata", "पार्वती माता", "Par-va-ti Ma-ta"),
        ("Kartikeya", "कार्तिकेय", "Kar-ti-ke-ya"),
        ("Mooshika", "मूषिक", "Moo-shi-ka"),
        ("Sanskrit", "संस्कृत", "Sans-krit"),
        ("Ganesh Chaturthi", "गणेश चतुर्थी", "Ga-nesh Cha-tur-thi"),
        ("Diwali", "दीवाली", "Di-wa-li"),
        ("Holi", "होली", "Ho-li"),
        ("Navratri", "नवरात्रि", "Nav-ra-tri"),
        ("Durga Puja", "दुर्गा पूजा", "Dur-ga Pu-ja"),
        ("Dussehra", "दशहरा", "Dus-seh-ra"),
        ("Raksha Bandhan", "रक्षा बंधन", "Rak-sha Band-han"),
        ("Janmashtami", "जन्माष्टमी", "Jan-mash-ta-mi"),
        ("Pongal", "पोंगल", "Pon-gal"),
        ("Onam", "ओणम", "O-nam"),
    ]),
    ("Symbol And Story Words", [
        ("Modak", "मोदक", "Mo-dak"),
        ("Lotus", "कमल", "Lo-tus / Ka-mal"),
        ("Trunk", "सूंड", "Trunk / Soond"),
        ("Tusk", "दंत", "Tusk / Dant"),
        ("Big Belly", "लंबोदर", "Big Bel-ly / Lam-bo-dar"),
        ("Eyes", "आंखें", "Eyes / Aan-khen"),
        ("Ears", "कान", "Ears / Kaan"),
    ]),
]


SCENES = [
    ("Current Active Sources Checked", [
        ("Use this doc instead of older VO extract files", [
            "Active routing checked in App.jsx and sceneRegistry.js.",
            "Shared VO checked in src/lib/config/content/voiceGuidance.js.",
            "Opening modal VO checked in src/lib/config/content/openingModals.js.",
            "Scene-local child games checked for Modak, Pond, Symbol/Tusk, Shloka River, and Lambodara Lodge.",
            "Recently changed VO included: Almost there!, Look for another part of the shadow., Keep holding., Who can help tie the rope?, Drag the next log.",
        ]),
    ]),
    ("Opening Navigation And Map", [
        ("First Screens", [
            "Tap to continue.",
            "Tap anywhere to begin.",
            "Hi bestie... I'm Ganesha. Come play with me.",
            "Tap Start, and we'll begin our adventure.",
        ]),
        ("Profile And Story", [
            "What should I call you?",
            "Tell me your name.",
            "How old are you?",
            "Pick your friend!",
            "Nice choice!",
            "Yay, {childName}! Let's go!",
            "Are you ready? Let's meet Ganesha!",
            "My mom Parvati made me with love and brought me to life!",
            "Mom said, Guard the door! But uh-oh, the visitor was Dad Shiva!",
            "Mom felt very sad, so Dad gave me a magical elephant head!",
            "Now we were together again, as one happy family!",
            "And now, let's explore my world together!",
        ]),
        ("Map", [
            "Come - let's start at Modak Mountain.",
            "The Shloka River is waiting for you.",
            "Lambodara Lodge is waiting for you.",
            "Wonderful! You discovered all the symbols of Symbol Mountain.",
            "Beautiful chanting! You completed the Shloka River.",
            "I loved learning about you! You completed Lambodara Lodge.",
        ]),
    ]),
    ("Modak Mountain Scene 1 - Modak", [
        ("Opening And Calm Mooshika", [
            "Mooshika has a journey ahead. Let's help him.",
            "Mooshika's thoughts keep jumping around. Press and hold to help him focus.",
            "Press and hold Mooshika.",
            "Keep holding until Mooshika settles.",
            "Keep holding.",
            "Mooshika stayed with it and made a little progress.",
            "I can keep going, one small step at a time.",
        ]),
        ("Flower Journey", [
            "Two flowers are hiding in the leaves. Swipe them apart.",
            "There they are! Two flowers.",
            "Swipe the branches apart.",
            "Swipe all the way to open the branches.",
            "Try swiping sideways.",
            "Swipe sideways.",
            "Swipe a little farther.",
            "Swipe farther.",
            "Two flowers are too high for Mooshika. Pull the branch down.",
            "Pull the branch down and hold it.",
            "Pull the branch down and keep holding.",
            "Got them! Two more flowers.",
            "Pull it downward.",
            "Pull downward.",
            "Keep pulling down.",
            "The last two flowers are across the mud. Guide Mooshika from safe spot to safe spot.",
            "Guide Mooshika to the next safe spot.",
            "Drag Mooshika to the glowing safe spot.",
            "Good jump! Now the next safe spot.",
            "Ooh... that's wobbly.",
            "Back to the safe spot. Try again.",
            "You made it! All six flowers.",
        ]),
        ("Garland And Modak", [
            "You found all six flowers. Tap them to make a garland for Ganesha.",
            "Tap any flower.",
            "Tap this flower.",
            "The garland is ready! Drag it to Ganesha.",
            "Drag the garland to Ganesha.",
            "Drag the garland onto Ganesha.",
            "You brought the garland to Ganesha - and look, a sweet Modak!",
            "Mooshika felt upset, sad, nervous, pleased, happy, and excited.",
            "His big belly held all those feelings.",
            "I can have lots of feelings and keep going.",
            "I can learn, keep trying, and find joy.",
            "Small steps, big feelings, and lots of trying brought Mooshika to Ganesha.",
        ]),
    ]),
    ("Modak Mountain Scene 2 - Pond", [
        ("Full Scene", [
            "The lotus needs water to bloom. Let's help the water reach it.",
            "The rock is blocking the stream. Press and hold to move it.",
            "Keep holding.",
            "Press and hold the rock.",
            "Keep holding until the rock moves.",
            "The flowers are in the water's way. Gently move them aside.",
            "Gently move the flowers aside.",
            "Drag the flowers to the side.",
            "The water reached the lotus.",
            "Hold the bud to help it bloom.",
            "Press and hold the bud.",
            "Keep holding until it blooms.",
            "The lotus bloomed, even in the muddy pond.",
            "You used strength, then gentleness, and helped the lotus bloom.",
        ]),
    ]),
    ("Modak Mountain Scene 3 - Symbol / Forest Friends", [
        ("Scene-Level Lines", [
            "The forest friends need your help. Let's see what they need.",
            "Some things are hiding. Look closely and tap each one you find.",
            "Our friends need the right sounds. Listen, then tap where each sound came from.",
            "You looked closely and found them all.",
            "You listened carefully and found every sound.",
            "Look closely - there's still one hiding.",
            "Listen again - where did it come from?",
            "You looked, listened, and used what you had to help.",
        ]),
        ("Eyes Game", [
            "Monkey is hungry, and Peacock has lost a special feather. Look carefully. Can you find what they need?",
            "You looked closely and found them both!",
            "That is interesting, but it is not what they are looking for.",
            "Feather found! Peacock will be so happy.",
            "Mango found! Monkey is hungry no more.",
            "Now find the feather.",
            "Now find the mango.",
            "Look near the bushes. Something is peeking out.",
            "Tap the glowing bushes.",
        ]),
        ("Ears Game", [
            "Listen closely. Hear each place first. Then choose the sound that matters.",
            "Listen for the water.",
            "Listen for the water drip.",
            "Listen for the grass.",
            "Listen for the grass rustle.",
            "Now choose the sound.",
            "Listen again.",
            "Tap the glowing place.",
            "Elephant found water.",
            "Cow found grass.",
            "You listened carefully and found what mattered.",
        ]),
    ]),
    ("Modak Mountain Scene 4 - Sacred Assembly", [
        ("Opening And Cards", [
            "You found all eight symbols. Let's see what you remember.",
            "Look at the symbol, then tap where it belongs.",
            "Eyes. I can focus on what matters.",
            "Ears. I can listen for what helps me.",
            "Trunk. I can be strong and gentle.",
            "Tusk. I can use what I have for something important.",
            "Modak. I can learn, keep trying, and find joy.",
            "Lotus. I can keep growing when things are hard.",
            "Big Belly. I can have lots of feelings and keep going.",
            "Mooshika. I can keep going, one small step at a time.",
            "Try another place.",
            "You're remembering what Ganesha's symbols mean.",
            "Almost there - six symbols are back with Ganesha.",
            "You remembered all eight symbols and what they mean.",
        ]),
    ]),
    ("Shloka River Scene 1 - Vakratunda Grove", [
        ("Vakratunda And Mahakaya", [
            "Our river friends need a way across. Let's help.",
            "Round 1 done!",
            "Round 2 done!",
            "Lotus blooming!",
            "I can change my plan and try another way.",
            "Mahakaya - start!",
            "Mahakaya round 1 done!",
            "Mahakaya round 2 done!",
            "Amazing!",
            "I can make space in my heart for others.",
            "Frog's way to his family is blocked. Drag Frog to find another way.",
            "That way is blocked. Can you find another way around?",
            "Vakratunda! You found another way and helped Frog reach his family.",
            "Try going around it.",
            "Everyone wants to cross the river. Drag the logs to build a raft.",
            "Drag the log into the water.",
            "Put the log here.",
            "Drag the next log.",
            "Keep adding logs to the raft.",
            "Mahakaay! You made space so everyone could come along.",
            "Wonderful!",
            "Mahakaay - great form!",
            "You found another way, and made sure everyone could come along.",
        ]),
    ]),
    ("Shloka River Scene 2 - Suryakoti Bank", [
        ("Suryakoti And Samaprabha", [
            "The river is dark. Let's bring back some light.",
            "Rub the glowing spot.",
            "Keep rubbing.",
            "Almost there!",
            "Keep rubbing to bring back more light.",
            "Bunny can't find home. Rub the glow to bring back the light.",
            "Suryakoti! You brought back the light and helped Bunny home.",
            "Press and hold on the shadow.",
            "Press and hold another part of the shadow.",
            "Look for another part of the shadow.",
            "Keep holding.",
            "Deer sees a strange shadow. Press and hold to look closely.",
            "Samaprabha! You looked closely and helped Deer feel safe.",
            "You brought back the light and saw what was really there.",
        ]),
    ]),
    ("Shloka River Scene 3 - Nirvighnam Chant", [
        ("Nirvighnam And Kuru Me Deva", [
            "The way ahead is tricky. Let's help our friends through.",
            "Turtle's way home is blocked. Press and hold the rock to move it.",
            "The branch is still blocking the way. Drag it onto the bank.",
            "The reeds are blocking the way. Swipe them apart.",
            "Drag the branch onto the bank.",
            "Swipe the reeds apart.",
            "Baby Beaver is across the river.",
            "You cleared the way and helped Turtle through.",
            "The logs are too heavy. Drag the help bubble to someone who can help.",
            "The rope needs tying. Drag the help bubble to someone who can help.",
            "Elephant helped put the logs in place.",
            "Monkey helped tie the rope.",
            "Now drag the planks into place.",
            "Good. Continue dragging the planks.",
            "Drag the help bubble to someone who can help.",
            "Drag the help bubble to Elephant.",
            "Drag the help bubble to Monkey.",
            "Who can help tie the rope?",
            "Drag a plank toward the bridge.",
            "Drag the plank onto the glowing place.",
            "Kuru Me Deva! You asked for help and built the bridge.",
            "You cleared the path and asked for help when things got hard.",
        ]),
    ]),
    ("Shloka River Scene 4 - Sarvakaryeshu Chant", [
        ("Sarvakaryeshu And Sarvada", [
            "Let's see where Ganesha is with you today.",
            "So much to do. Which symbol can help him start small?",
            "Mooshika helped him begin with one small thing.",
            "So many distractions. Which symbol can help him focus?",
            "The Eyes helped him focus on what mattered.",
            "Heavy here, gentle there. Which symbol can help?",
            "The Trunk helped him use strength, then gentleness.",
            "The crayon is broken. Which symbol can help?",
            "The Tusk reminded her that something broken can still help.",
            "Which symbol can help him start small?",
            "Think of a symbol that reminds us to begin with one small step.",
            "Remember? Mooshika helps us start small.",
            "Which symbol helps us be strong and gentle?",
            "Think of a symbol that can be strong and gentle.",
            "Remember? The Trunk can be strong and gentle.",
            "Which symbol can help him focus?",
            "Think of the symbol that helps us notice what matters.",
            "Remember? The Eyes help us focus on what matters.",
            "Which symbol reminds us broken things can still help?",
            "Think of a symbol that shows something broken can still be useful.",
            "Remember? The Tusk can still help, even when it's broken.",
            "Morning, afternoon and night - let's find Ganesha through your day.",
            "Morning can be busy. What can help him listen?",
            "Things didn't go his way. What can help him keep going?",
            "He kept trying and finished the puzzle. What can remind him of that happy feeling?",
            "Can you find the hidden symbol?",
            "Look for something that helps us listen.",
            "Look closely at the backpack.",
            "Look for something that can hold lots of feelings.",
            "Look at the playground.",
            "Look for something sweet that reminds us of joy.",
            "Look near the lamp.",
            "Look through the picture and find Ganesha's hidden symbol.",
            "You found Ganesha through the whole day.",
            "From morning to night, Ganesha was with you.",
        ]),
    ]),
    ("Shloka River Scene 5 - Finale", [
        ("Full Shloka Review", [
            "You found all eight words. Now let's put them together.",
            "Watch the Shloka come alive.",
            "Now, choose the first symbol boat in the Shloka.",
            "Now, choose the next one.",
            "Great! Keep choosing.",
            "You did it! The whole Shloka is in place.",
            "Now listen to the whole shloka.",
            "Now, let's say it together.",
            "You remembered the whole Ganesha Shloka. Now you can carry it with you.",
            "Tap the glowing boat.",
            "Remember Ganesha's curved trunk.",
            "Think of Ganesha's great form.",
            "Think of a light as bright as millions of suns.",
            "Think of steady, equal brightness.",
            "Remember the word about clearing obstacles.",
            "Remember the words that ask for help.",
            "Think of the word for everything we do.",
            "Remember the word that means always.",
            "Which symbol can help you try another way?",
            "Remember Vakratunda.",
            "Which symbol reminds you to make space for others?",
            "Remember Mahakaay.",
            "Which symbol can help you find hope inside?",
            "Remember Suryakoti.",
            "Which symbol can help you see what is really there?",
            "Remember Samaprabha.",
            "Which symbol can help you solve one thing at a time?",
            "Remember Nirvighnam.",
            "Which symbol reminds you to ask for help?",
            "Remember Kuru Me Deva.",
            "Which symbol reminds you of Ganesha in all you do?",
            "Remember Sarva Karyeshu.",
            "Which symbol reminds you Ganesha is always with you?",
            "Remember Sarvada.",
        ]),
    ]),
    ("Lambodara Lodge Scene 1 - My Indian Story", [
        ("India Story", [
            "Let's explore my India story, then yours.",
            "Move the magnifying glass around the map to find my places.",
            "Now choose a place your family feels connected to.",
            "Which place feels connected to your family?",
            "This is the place you chose.",
            "Tap play and listen carefully.",
            "Which language do you hear?",
            "Think of the language many shlokas are in.",
            "Vakratunda Mahakaya Suryakoti Samaprabha!",
            "Yes - that's Sanskrit, a language used in many mantras.",
            "Now choose up to three languages you use or hear.",
            "These are the languages you chose.",
            "Choose the languages that are part of your life.",
            "Can you find my festival?",
            "Think of the festival that celebrates Ganesha.",
            "Yes - Ganesh Chaturthi!",
            "Now choose the festivals your family celebrates.",
            "Which celebrations are part of your family?",
            "These are the festivals you chose.",
            "Your places, languages, and celebrations can all be part of that connection.",
            "You found some places connected to my story.",
        ]),
    ]),
    ("Lambodara Lodge Scene 2 - Family Tree", [
        ("Family Tree", [
            "Let's meet my family, then make your family tree.",
            "This is my father, Shiva.",
            "This is my mother, Parvati.",
            "This is my brother, Kartikeya.",
            "And that's me!",
            "Our family trees are different, and each one tells a story.",
        ]),
    ]),
    ("Lambodara Lodge Scene 3 - Favorite Things", [
        ("Favorite Things", [
            "Let's find my favorites, then yours.",
            "Can you guess my favorite food?",
            "Look for the sweet I love.",
            "Yes - I love modak!",
            "Can you guess my favorite color?",
            "What color is an apple?",
            "Yes - red!",
            "What do you think I love doing?",
            "What lets me move to music?",
            "Yes - I love to dance!",
            "Who do you think my friend is?",
            "Look for my tiny friend.",
            "Yes - Mooshika!",
            "Now I want to know your favorites.",
            "Choose your favorite food.",
            "Yummy!",
            "Choose your favorite color.",
            "Nice choice!",
            "Choose something you love to do.",
            "That sounds like fun!",
            "Add the name of a friend you care about.",
            "Thanks for telling me!",
            "Now we know some things we each enjoy.",
        ]),
    ]),
    ("Lambodara Lodge Scene 4 - Wishes And Dreams", [
        ("Wishes And Dreams", [
            "I have a few wishes. Will you help me - then share yours?",
            "My first wish is for more kindness.",
            "Tap the kind actions you can spot.",
            "Tap the bubbles that show helping, sharing, hugging, and gifting.",
            "Look for actions that help or care for someone.",
            "Oh no, that's not kind. Look for a helping bubble.",
            "Can you find another kind action?",
            "You found all the kind choices.",
            "My next wish is for everyone to have enough food.",
            "Drag food to each empty plate.",
            "There's still an empty plate.",
            "Now every plate has food.",
            "My next wish is for a greener world.",
            "Tap the bare spaces to help the garden grow.",
            "There's still space for more to grow.",
            "Look how green it has become.",
            "Now it's your turn to make a wish.",
            "Draw something you hope for or dream about.",
            "Your dream is hiding behind the clouds. Help me clear them - tap me.",
            "There are still clouds in the way.",
            "There it is - your dream!",
            "Everyone can dream about something different.",
        ]),
    ]),
]


def set_run(run, size=10.5, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, 20, True, (85, 40, 115))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, 15 if level == 1 else 12, True, (85, 40, 115) if level == 1 else (80, 80, 80))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run(r, 10.5)


def add_word_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Word", "Hindi / Sanskrit", "Connected Recording Breakup"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(header)
        set_run(r, 10, True, (255, 255, 255))
        cell._tc.get_or_add_tcPr().append(parse_shd("5B2B7A"))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            r = cells[i].paragraphs[0].add_run(value)
            set_run(r, 9.5)
    return table


def parse_shd(fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    return shd


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    add_title(doc, "VO Artist Handoff - Current Active VO")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Hindi/Sanskrit connected words first, followed by scene-wise recording scripts.")
    set_run(r, 10.5, False, (80, 80, 80))

    add_heading(doc, "Recording Notes", 1)
    for note in [
        "Audience: young children. Keep the delivery warm, clear, playful, and calm.",
        "Keep Hindi/Sanskrit words clean and unhurried. Do not over-anglicize the connected words.",
        "Lines with {childName} or similar placeholders are dynamic. Record the surrounding phrase only if the production plan needs it.",
        "This version is built from current active scene files and shared VO config, not the older extraction notes.",
    ]:
        add_bullet(doc, note)

    doc.add_page_break()
    add_heading(doc, "Hindi / Sanskrit Connected Words", 1)
    for section_title, rows in WORD_SECTIONS:
        add_heading(doc, section_title, 2)
        add_word_table(doc, rows)
        doc.add_paragraph()

    doc.add_page_break()
    add_heading(doc, "Scene-Wise VO Scripts", 1)
    for scene_title, blocks in SCENES:
        add_heading(doc, scene_title, 1)
        for block_title, lines in blocks:
            add_heading(doc, block_title, 2)
            for line in lines:
                add_bullet(doc, line)
        doc.add_paragraph()

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
